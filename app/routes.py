import hmac
import io
import logging
from datetime import datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo
from flask import Blueprint, current_app, jsonify, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from .extensions import db, limiter
from .models import Customer, Prize, Spin
from .services import create_spin
from .utils import customer_dict, error, normalize_mobile, prize_dict, spin_dict

api = Blueprint('api', __name__, url_prefix='/api')
logger = logging.getLogger(__name__)


def payload():
    return request.get_json(silent=True) if request.is_json else None


def today_completed_rows():
    """Return completed spins for the store's current India calendar day."""
    india = ZoneInfo('Asia/Kolkata')
    today = datetime.now(india).date()
    start = datetime.combine(today, time.min, tzinfo=india).astimezone(timezone.utc)
    end = start + timedelta(days=1)
    rows = db.session.execute(
        select(Customer, Spin, Prize)
        .join(Spin, Spin.customer_id == Customer.id)
        .join(Prize, Prize.id == Spin.prize_id)
        .where(Spin.spin_time >= start, Spin.spin_time < end, Spin.status == 'completed')
        .order_by(Spin.spin_time.asc())
    ).all()
    return today, india, rows


def current_customer_response(customer):
    spin = customer.spin
    return {
        'success': True, 'exists': True, 'can_spin': not bool(spin), 'customer': customer_dict(customer),
        'spin': spin_dict(spin) if spin else None,
    }


@api.get('/health')
def health():
    try:
        db.session.execute(text('SELECT 1'))
        return jsonify({'success': True, 'status': 'ok'})
    except SQLAlchemyError:
        logger.exception('Health database check failed')
        return error('Database is unavailable.', 503)


@api.post('/customer/check')
@limiter.limit('60 per minute')
def check_customer():
    data = payload() or {}
    mobile = normalize_mobile(data.get('mobile'))
    if not mobile:
        return error('Please enter a valid 10-digit Indian mobile number.')
    customer = db.session.scalar(select(Customer).where(Customer.mobile_number == mobile))
    if not customer:
        return jsonify({'success': True, 'exists': False, 'can_spin': True})
    return jsonify(current_customer_response(customer))


@api.post('/customer/register')
@limiter.limit('30 per minute')
def register_customer():
    data = payload() or {}
    mobile = normalize_mobile(data.get('mobile'))
    name = data.get('name', '').strip() if isinstance(data.get('name'), str) else ''
    address = data.get('address', '').strip() if isinstance(data.get('address'), str) else ''
    if not mobile:
        return error('Please enter a valid 10-digit Indian mobile number.')
    if not name or len(name) > 120:
        return error('Please provide a valid name.')
    if not address or len(address) > 2000:
        return error('Please provide a valid address.')
    customer = Customer(mobile_number=mobile, name=name, address=address)
    try:
        db.session.add(customer)
        db.session.commit()
        logger.info('Customer registered')
        return jsonify({'success': True, 'customer': customer_dict(customer), 'exists': False}), 201
    except IntegrityError:
        db.session.rollback()
        existing = db.session.scalar(select(Customer).where(Customer.mobile_number == mobile))
        if existing:
            return jsonify({**current_customer_response(existing), 'exists': True}), 200
        return error('Unable to save customer details. Please try again.', 409)
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception('Customer registration database error')
        return error('Unable to save your details right now. Please try again.', 503)


@api.get('/customer/<mobile>')
def get_customer(mobile):
    normalized = normalize_mobile(mobile)
    if not normalized:
        return error('Invalid mobile number.')
    customer = db.session.scalar(select(Customer).where(Customer.mobile_number == normalized))
    if not customer:
        return error('Customer not found.', 404)
    return jsonify(current_customer_response(customer))


@api.post('/customer/social')
@limiter.limit('30 per minute')
def complete_social():
    data = payload() or {}
    mobile = normalize_mobile(data.get('mobile'))
    completed = data.get('completed')
    if not mobile:
        return error('Please enter a valid 10-digit Indian mobile number.')
    if not isinstance(completed, dict) or not all(completed.get(key) is True for key in ('facebook', 'instagram', 'whatsapp')):
        return error('Please complete all social media requirements.')
    customer = db.session.scalar(select(Customer).where(Customer.mobile_number == mobile))
    if not customer:
        return error('Please complete your details first.', 404)
    customer.facebook_completed = True
    customer.instagram_completed = True
    customer.whatsapp_completed = True
    try:
        db.session.commit()
        return jsonify({'success': True, 'customer': customer_dict(customer)})
    except SQLAlchemyError:
        db.session.rollback()
        logger.exception('Social completion database error')
        return error('Unable to save your social media confirmation. Please try again.', 503)


@api.post('/spin')
@limiter.limit('20 per minute')
def spin():
    data = payload() or {}
    mobile = normalize_mobile(data.get('mobile'))
    if not mobile:
        return error('Please enter a valid 10-digit Indian mobile number.')
    try:
        customer, spin_record, outcome = create_spin(mobile)
    except SQLAlchemyError:
        logger.exception('Spin transaction database error')
        return error('Unable to complete your spin right now. Please try again.', 503)
    if outcome == 'missing':
        return error('Please complete your details before spinning.', 404)
    if outcome == 'social_incomplete':
        return error('Please complete all social media requirements before spinning.', 403)
    if outcome == 'error' or not spin_record:
        logger.error('Spin transaction did not create or find a spin')
        return error('Unable to complete your spin right now. Please try again.', 503)
    response = {'success': True, 'customer': customer_dict(customer), 'spin': spin_dict(spin_record), 'already_spun': outcome == 'existing'}
    logger.info('Existing spin returned' if outcome == 'existing' else 'Spin completed')
    return jsonify(response)


@api.get('/prizes')
def prizes():
    rows = db.session.scalars(select(Prize).where(Prize.is_active.is_(True))).all()
    return jsonify({'success': True, 'prizes': [prize_dict(prize) | {'weight': prize.weight} for prize in rows]})


@api.get('/social')
def social():
    return jsonify({'success': True, 'socials': [
        {'name': 'Instagram', 'url': 'https://www.instagram.com/mobilehub_tadepalligudem_/'},
        {'name': 'Facebook', 'url': 'https://www.facebook.com/profile.php?id=61555352783316'},
        {'name': 'WhatsApp', 'url': 'https://whatsapp.com/channel/0029VajCJxRADTOB8x4ocV2t'},
    ]})


@api.get('/store')
def store():
    return jsonify({'success': True, 'name': 'Mobile Hub', 'address': ['Opp Prabhata Talkies', 'Bhimavaram Road', 'Tadepalligudem - 534102'], 'phone': '9014 567 567'})


@api.post('/admin/login')
@limiter.limit('10 per minute')
def admin_login():
    data = payload() or {}
    password = data.get('password', '')
    if not isinstance(password, str) or not hmac.compare_digest(password, current_app.config['ADMIN_PASSWORD']):
        return error('Incorrect admin password.', 401)
    session.clear()
    session['admin_authenticated'] = True
    return jsonify({'success': True})


@api.get('/admin/session')
def admin_session():
    return jsonify({'success': True, 'authenticated': bool(session.get('admin_authenticated'))})


@api.get('/admin/entries')
def admin_today_entries():
    if not session.get('admin_authenticated'):
        return error('Admin login is required.', 401)
    _today, india, rows = today_completed_rows()
    entries = [
        {
            'customer_name': customer.name,
            'mobile_no': customer.mobile_number,
            'address': customer.address,
            'gift': prize.title,
            'created_at': customer.created_at.astimezone(india).isoformat(),
        }
        for customer, spin_record, prize in rows
    ]
    return jsonify({'success': True, 'count': len(entries), 'entries': entries})


@api.post('/export/customers')
def export_customers():
    """Export every real customer record, including a reward when one exists."""
    secret = request.headers.get('X-Export-Secret', '')
    configured = current_app.config['EXPORT_SECRET']
    has_export_secret = bool(configured) and hmac.compare_digest(secret, configured)
    if not session.get('admin_authenticated') and not has_export_secret:
        return error('Export authorization failed.', 401)
    # `customers` is the live MySQL customer table. Left joins keep customers
    # who have not spun yet in the export, with their reward marked “Not spun”.
    rows = db.session.execute(
        select(Customer, Spin, Prize)
        .outerjoin(Spin, Spin.customer_id == Customer.id)
        .outerjoin(Prize, Prize.id == Spin.prize_id)
        .order_by(Customer.created_at.desc())
    ).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Customers'
    sheet.append(['Customer Name', 'Mobile Number', 'Address', 'Facebook', 'Instagram', 'WhatsApp', 'Reward Won', 'Date', 'Time'])
    header_fill = PatternFill('solid', fgColor='771018')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for customer, spin_record, prize in rows:
        event_time = spin_record.spin_time if spin_record else customer.created_at
        sheet.append([customer.name, customer.mobile_number, customer.address, 'Completed' if customer.facebook_completed else 'Pending', 'Completed' if customer.instagram_completed else 'Pending', 'Completed' if customer.whatsapp_completed else 'Pending', prize.title if prize else 'Not spun', event_time.astimezone(timezone.utc).strftime('%Y-%m-%d'), event_time.astimezone(timezone.utc).strftime('%H:%M:%S UTC')])
    for column, width in zip('ABCDEFGHI', (26, 16, 40, 14, 14, 14, 38, 14, 18)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column == 3)
    output = io.BytesIO(); workbook.save(output); output.seek(0)
    filename = f"mobile-hub-spin-win-{datetime.now(timezone.utc).date().isoformat()}.xlsx"
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@api.post('/export/today')
def export_today():
    """Download today's completed spins using the store's local calendar day."""
    if not session.get('admin_authenticated'):
        return error('Export authorization failed.', 401)

    today, _india, rows = today_completed_rows()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Today Entries'
    sheet.append(['customer_name', 'mobile_no', 'address', 'gift'])
    for customer, _spin_record, prize in rows:
        sheet.append([customer.name, customer.mobile_number, customer.address, prize.title])
    for column, width in zip('ABCD', (26, 16, 40, 38)):
        sheet.column_dimensions[column].width = width
    output = io.BytesIO(); workbook.save(output); output.seek(0)
    filename = f'MobileHub_Today_Entries_{today.isoformat()}.xlsx'
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
